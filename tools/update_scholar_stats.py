#!/usr/bin/env python3
"""Record one daily Google Scholar profile snapshot in an Excel workbook."""

from __future__ import annotations

import json
import os
import re
from datetime import datetime, timedelta, timezone
from html import unescape
from pathlib import Path
from urllib.parse import parse_qs, urlparse
from urllib.request import Request, urlopen

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent.parent
PROFILE_URL = "https://scholar.google.com/citations?user=PyRqpAsAAAAJ&hl=en&pagesize=100"
EXCEL_FILE = ROOT / "assets" / "data" / "scholar-citations.xlsx"
SITE_DATA_FILE = ROOT / "assets" / "data" / "scholar.json"
CHINA_TIME = timezone(timedelta(hours=8))
HEADERS = ("日期（北京时间）", "总引用", "较上次记录变化", "H 指数", "i10 指数", "获取时间（北京时间）")
PAPER_HEADERS = ("日期（北京时间）", "论文 ID", "论文标题", "引用数", "较该论文上次记录变化", "获取时间（北京时间）")


def parse_profile(source: str) -> tuple[dict[str, int], dict[str, dict]]:
    if "captcha" in source.lower() or "not a robot" in source.lower():
        raise ValueError("Google Scholar requested human verification")
    values = re.findall(r'class="gsc_rsb_std"[^>]*>\s*([\d,]+)', source)
    if len(values) < 5:
        raise ValueError("Google Scholar profile metrics are missing")
    metrics = {
        "citations": int(values[0].replace(",", "")),
        "hindex": int(values[2].replace(",", "")),
        "i10index": int(values[4].replace(",", "")),
    }
    papers = {}
    rows = re.findall(r'<tr\b[^>]*class="[^"]*\bgsc_a_tr\b[^"]*"[^>]*>(.*?)</tr>', source, re.S)
    for row in rows:
        title_match = re.search(r'<a\b([^>]*class="[^"]*\bgsc_a_at\b[^"]*"[^>]*)>(.*?)</a>', row, re.S)
        if not title_match:
            continue
        href_match = re.search(r'href="([^"]+)"', title_match.group(1))
        if not href_match:
            continue
        paper_id = parse_qs(urlparse(unescape(href_match.group(1))).query).get("citation_for_view", [None])[0]
        title = unescape(re.sub(r'<[^>]+>', '', title_match.group(2))).strip()
        count_match = re.search(r'<a\b[^>]*class="[^"]*\bgsc_a_ac\b[^"]*"[^>]*>(.*?)</a>', row, re.S)
        count_text = re.sub(r'<[^>]+>', '', count_match.group(1)).strip() if count_match else ''
        if not paper_id or not title or (count_text and not re.fullmatch(r'[\d,]+', count_text)):
            continue
        papers[paper_id] = {"title": title, "citations": int(count_text.replace(',', '')) if count_text else 0}
    if not rows or len(papers) != len(rows):
        raise ValueError("Google Scholar publication list is missing or incomplete")
    return metrics, papers


def format_header(sheet, headers: tuple[str, ...], widths: tuple[int, ...]) -> None:
    sheet.append(headers)
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    for index, width in enumerate(widths, start=1):
        cell = sheet.cell(1, index)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243B53")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 26


def create_workbook() -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "每日引用"
    format_header(sheet, HEADERS, (18, 16, 20, 14, 14, 25))
    format_header(workbook.create_sheet("逐篇论文"), PAPER_HEADERS, (18, 28, 90, 16, 24, 25))
    return workbook


def record_snapshot(metrics: dict[str, int], papers: dict[str, dict], observed_at: datetime) -> None:
    local_time = observed_at.astimezone(CHINA_TIME).replace(tzinfo=None)
    today = local_time.date()
    workbook = load_workbook(EXCEL_FILE) if EXCEL_FILE.exists() else create_workbook()
    sheet = workbook["每日引用"]
    if tuple(cell.value for cell in sheet[1]) != HEADERS:
        raise ValueError("Unexpected Excel header; existing workbook was left unchanged")
    existing_row = next((row for row in range(2, sheet.max_row + 1)
                         if sheet.cell(row, 1).value == datetime.combine(today, datetime.min.time())), None)
    row = existing_row or sheet.max_row + 1
    sheet.cell(row, 1, datetime.combine(today, datetime.min.time()))
    sheet.cell(row, 2, metrics["citations"])
    sheet.cell(row, 4, metrics["hindex"])
    sheet.cell(row, 5, metrics["i10index"])
    sheet.cell(row, 6, local_time)
    sheet.cell(row, 1).number_format = "yyyy-mm-dd"
    sheet.cell(row, 6).number_format = "yyyy-mm-dd hh:mm"
    for col in (2, 3, 4, 5):
        sheet.cell(row, col).number_format = "#,##0;[Red]-#,##0;0"
    for col in range(1, 7):
        sheet.cell(row, col).font = Font(name="Arial", size=10, color="243B53")
        sheet.cell(row, col).alignment = Alignment(vertical="center", horizontal="right" if col > 1 else "center")
    sheet.row_dimensions[row].height = 22
    for current_row in range(2, sheet.max_row + 1):
        previous = sheet.cell(current_row - 1, 2).value if current_row > 2 else None
        current = sheet.cell(current_row, 2).value
        sheet.cell(current_row, 3, current - previous if previous is not None else None)
    sheet.auto_filter.ref = f"A1:F{sheet.max_row}"
    detail = workbook["逐篇论文"] if "逐篇论文" in workbook.sheetnames else workbook.create_sheet("逐篇论文")
    if detail.max_row == 1 and detail.cell(1, 1).value is None:
        format_header(detail, PAPER_HEADERS, (18, 28, 90, 16, 24, 25))
    if tuple(cell.value for cell in detail[1]) != PAPER_HEADERS:
        raise ValueError("Unexpected paper sheet header; existing workbook was left unchanged")
    previous_counts = {}
    for detail_row in range(2, detail.max_row + 1):
        recorded_date = detail.cell(detail_row, 1).value
        paper_id = detail.cell(detail_row, 2).value
        if recorded_date and recorded_date.date() < today and paper_id:
            previous_counts[paper_id] = detail.cell(detail_row, 4).value
    for detail_row in range(detail.max_row, 1, -1):
        recorded_date = detail.cell(detail_row, 1).value
        if recorded_date and recorded_date.date() == today:
            detail.delete_rows(detail_row)
    for paper_id, paper in sorted(papers.items(), key=lambda item: item[1]["title"].casefold()):
        count = paper["citations"]
        previous = previous_counts.get(paper_id)
        detail.append((datetime.combine(today, datetime.min.time()), paper_id, paper["title"],
                       count, count - previous if previous is not None else None, local_time))
        detail_row = detail.max_row
        detail.cell(detail_row, 1).number_format = "yyyy-mm-dd"
        detail.cell(detail_row, 6).number_format = "yyyy-mm-dd hh:mm"
        for col in (4, 5):
            detail.cell(detail_row, col).number_format = "#,##0;[Red]-#,##0;0"
        for col in range(1, 7):
            detail.cell(detail_row, col).font = Font(name="Arial", size=10, color="243B53")
            detail.cell(detail_row, col).alignment = Alignment(vertical="center", horizontal="left" if col in (2, 3) else "right")
        detail.cell(detail_row, 3).alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
        detail.row_dimensions[detail_row].height = 42
    detail.auto_filter.ref = f"A1:F{detail.max_row}"
    EXCEL_FILE.parent.mkdir(parents=True, exist_ok=True)
    temporary = EXCEL_FILE.with_suffix(".tmp.xlsx")
    try:
        workbook.save(temporary)
        os.replace(temporary, EXCEL_FILE)
    finally:
        if temporary.exists():
            temporary.unlink()
    # Keep the small existing website cache current; Excel is the only history file.
    SITE_DATA_FILE.write_text(json.dumps({**metrics, "updated": observed_at.replace(microsecond=0).isoformat().replace("+00:00", "Z")},
                                         ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Recorded {today}: {metrics['citations']} citations across {len(papers)} papers")


def main() -> None:
    request = Request(PROFILE_URL, headers={"User-Agent": "Mozilla/5.0 (compatible; academic-homepage-updater/1.0)"})
    try:
        with urlopen(request, timeout=30) as response:
            source = response.read().decode("utf-8", errors="replace")
        metrics, papers = parse_profile(source)
    except (OSError, ValueError) as error:
        print(f"Scholar unavailable; previous Excel data preserved: {error}")
        return
    record_snapshot(metrics, papers, datetime.now(timezone.utc))


if __name__ == "__main__":
    main()
